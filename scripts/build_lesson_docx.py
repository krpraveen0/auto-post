#!/usr/bin/env python3
"""Build a polished DOCX lesson from the repository's portable Markdown."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
SLATE = RGBColor(71, 85, 105)
BLUE = RGBColor(30, 64, 175)
LIGHT_BLUE = "EAF2FF"
LIGHT_GRAY = "F4F6F8"


def set_cell_shading(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_repeat_keep(paragraph, *, keep_next: bool = False) -> None:
    properties = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    properties.append(keep_lines)
    if keep_next:
        item = OxmlElement("w:keepNext")
        properties.append(item)


def remove_paragraph_border(paragraph_or_style) -> None:
    element = paragraph_or_style._element
    properties = element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "between"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1E40AF")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("["):
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link and re.match(r"https?://", link.group(2)):
                add_hyperlink(paragraph, link.group(1), link.group(2))
            elif link:
                paragraph.add_run(f"{link.group(1)} ({link.group(2)})")
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Courier New", 9.5)
            run.font.color.rgb = RGBColor(31, 41, 55)
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token[1:-1]).italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def set_picture_alt_text(run, alt_text: str) -> None:
    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        return
    doc_pr = drawing.find(".//" + qn("wp:docPr"))
    if doc_pr is not None:
        doc_pr.set("descr", alt_text)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(8)
    remove_paragraph_border(title)

    for name, size, before, after in (
        ("Heading 1", 18, 18, 7),
        ("Heading 2", 14, 14, 5),
        ("Heading 3", 12, 11, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Courier New"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor(31, 41, 55)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(2)
    code.paragraph_format.line_spacing = 1.0

    if "Caption" in styles:
        caption = styles["Caption"]
        caption.font.name = "Arial"
        caption.font.size = Pt(9)
        caption.font.italic = True
        caption.font.color.rgb = SLATE
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(10)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Agentic AI Engineering  |  Part 3                                      ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def markdown_body(path: Path) -> tuple[dict[str, str], list[str]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    metadata: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1)
        for line in lines[1:end]:
            if ":" in line:
                key, value = line.split(":", 1)
                metadata[key.strip()] = value.strip()
        lines = lines[end + 1 :]
    return metadata, lines


def add_code_block(document: Document, lines: list[str]) -> None:
    keep_group = len(lines) <= 10
    for index, line in enumerate(lines or [""]):
        paragraph = document.add_paragraph(style="Code Block")
        set_cell_shading(paragraph, LIGHT_GRAY)
        run = paragraph.add_run(line or " ")
        set_run_font(run, "Courier New", 8.5)
        paragraph.paragraph_format.keep_together = True
        if keep_group and index < len(lines) - 1:
            paragraph.paragraph_format.keep_with_next = True
        if index < len(lines) - 1:
            paragraph.paragraph_format.space_after = Pt(0)


def build(source: Path, output: Path) -> None:
    metadata, lines = markdown_body(source)
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    document.core_properties.title = metadata.get("title", source.stem)
    document.core_properties.subject = "Agentic AI Engineering Part 3"
    document.core_properties.author = metadata.get("author", "Praveen Kumar")
    document.core_properties.keywords = metadata.get("tags", "AI agents, tool use")
    document.core_properties.comments = "Generated from the canonical repository Markdown."

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.style = document.styles["Normal"]
    footer.paragraph_format.space_before = Pt(4)
    footer.paragraph_format.space_after = Pt(0)
    footer_run = footer.add_run()
    set_run_font(footer_run, "Arial", 8)
    add_page_number(footer)
    for run in footer.runs:
        run.font.color.rgb = SLATE
        set_run_font(run, "Arial", 8)

    index = 0
    code_lines: list[str] | None = None
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(value.strip() for value in paragraph_buffer).strip()
        paragraph_buffer = []
        if text:
            paragraph = document.add_paragraph()
            add_inline(paragraph, text)

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if code_lines is not None:
            if stripped == "```":
                add_code_block(document, code_lines)
                code_lines = None
            else:
                code_lines.append(line)
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            code_lines = []
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            flush_paragraph()
            target = (source.parent / image.group(2)).resolve()
            if target.suffix.lower() == ".svg":
                png_target = target.with_suffix(".png")
                if png_target.exists():
                    target = png_target
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(target), width=Inches(5.55))
            set_picture_alt_text(run, image.group(1))
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                remove_paragraph_border(paragraph)
                add_inline(paragraph, text)
                subtitle = metadata.get("subtitle")
                if subtitle:
                    sub = document.add_paragraph()
                    sub.paragraph_format.space_after = Pt(6)
                    run = sub.add_run(subtitle)
                    run.font.size = Pt(14)
                    run.font.color.rgb = SLATE
                meta = document.add_paragraph()
                meta.paragraph_format.space_after = Pt(12)
                meta_run = meta.add_run(
                    f"By {metadata.get('author', 'Praveen Kumar')}  |  Course lesson 3  |  Verified {metadata.get('last_verified', '')}"
                )
                meta_run.font.size = Pt(9)
                meta_run.font.bold = True
                meta_run.font.color.rgb = BLUE
            else:
                paragraph = document.add_paragraph(style=f"Heading {level - 1}")
                add_inline(paragraph, text)
                set_repeat_keep(paragraph, keep_next=True)
            index += 1
            continue
        caption = re.match(r"^\*\*(Figure\s+\d+:[^*]+)\*\*$", stripped)
        if caption:
            flush_paragraph()
            paragraph = document.add_paragraph(style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(caption.group(1))
            index += 1
            continue
        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if ordered or bullet:
            flush_paragraph()
            text = ordered.group(2) if ordered else bullet.group(1)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.24)
            prefix = f"{ordered.group(1)}. " if ordered else "•  "
            paragraph.add_run(prefix)
            add_inline(paragraph, text)
            paragraph.paragraph_format.space_after = Pt(3)
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    if code_lines is not None:
        add_code_block(document, code_lines)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
    print(f"Created {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
